"""Write small sample documents of every supported format for tests."""

import json
from pathlib import Path

import docx


def write_text_pdf(path: Path, pages: list[str]) -> Path:
    """Write a PDF whose pages contain extractable text (one line each).

    pypdf can only write blank pages, so the file is assembled by hand.
    """
    page_count = len(pages)
    font_id = 3 + 2 * page_count
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [{}] /Count {} >>".format(
            " ".join(f"{3 + 2 * n} 0 R" for n in range(page_count)), page_count
        ),
    ]

    for number, text in enumerate(pages):
        content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> "
            f"/Contents {4 + 2 * number} 0 R >>"
        )
        objects.append(f"<< /Length {len(content)} >>\nstream\n{content}\nendstream")

    objects.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    output = "%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output += f"{number} 0 obj\n{body}\nendobj\n"

    xref_offset = len(output)
    output += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n"
    output += "".join(f"{offset:010d} 00000 n \n" for offset in offsets)
    output += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n"
    )

    path.write_bytes(output.encode("latin-1"))
    return path


def write_docx(path: Path) -> Path:
    document = docx.Document()
    document.core_properties.author = "HR Team"
    document.core_properties.title = "Leave Policy"
    document.add_heading("Leave Policy", level=1)
    document.add_paragraph("Employees get 20 days of annual leave.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Type"
    table.cell(0, 1).text = "Days"
    table.cell(1, 0).text = "Sick"
    table.cell(1, 1).text = "10"
    document.add_paragraph("Unused leave carries over.")
    document.save(str(path))
    return path


def write_sample(directory: Path, extension: str) -> Path:
    """Write a sample file mentioning leave in the given format."""
    path = directory / f"sample{extension}"

    if extension == ".pdf":
        return write_text_pdf(path, ["Annual leave is 20 days."])
    if extension == ".docx":
        return write_docx(path)

    content = {
        ".txt": "Annual leave is 20 days.\n",
        ".md": "# Leave\n\nAnnual leave is **20 days**.\n",
        ".html": "<html><body><p>Annual leave is 20 days.</p></body></html>",
        ".csv": "type,days\nannual leave,20\n",
        ".json": json.dumps({"leave": {"annual": "20 days"}}),
    }[extension]
    path.write_text(content, encoding="utf-8")
    return path
